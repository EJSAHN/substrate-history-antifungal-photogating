"""Export analysis CSVs to the Supplementary Data 1 workbook.

Inferential statistics are computed in analysis.py. Display formulas are checked
against those results before their cached values are written.
"""
from __future__ import annotations
import argparse
import csv
import json
import math
import os
from pathlib import Path
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile


def column(n):
    s=''
    while n:n,k=divmod(n-1,26);s=chr(65+k)+s
    return s

def scalar(s):
    if s=='':return None
    if s in ['True','False']:return s=='True'
    try:
        v=float(s)
        if not math.isfinite(v):raise ValueError('Non-finite CSV value: '+s)
        return int(v) if v.is_integer() and abs(v)<2**53 else v
    except ValueError:
        if s.lower() in ['nan','inf','-inf','infinity']:raise
        return s

def source_csv(root,name,headers):
    hits=[root/d/name for d in ['tables','audit','legacy_reference','inputs_snapshot','source_data'] if (root/d/name).is_file()]
    if len(hits)>1:raise ValueError('Ambiguous result table: '+name)
    if not hits:
        optional={'Legacy_JMP_variable_clustering_NOT_recomputed.csv',
                  'historical_dose_CI_unresolved_comparison.csv','baseline_checks.csv'}
        if name in optional:return []
        raise FileNotFoundError('Required current-run table is missing: '+name)
    with hits[0].open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f)
        if not reader.fieldnames:return []
        if list(reader.fieldnames)!=list(headers):raise ValueError('Column schema changed for '+name)
        return [[scalar(r[h]) for h in headers] for r in reader]

def make_snapshot(root: Path):
    schema=json.loads((Path(__file__).resolve().parent/'schema/workbook_layout.json').read_text(encoding='utf8'))
    sheets=[];lookup={}
    def add(item,rows,formulas=None):
        n=len(item['headers'])
        values=[[item['title']]+[None]*(n-1),[item['description']]+[None]*(n-1),[None]*n,list(item['headers'])]+rows
        s={'name':item['name'],'columns':n,'data_rows':len(rows),'values':values,'formulas':formulas or [],'source':item.get('csv','Derived display / study metadata')}
        sheets.append(s);lookup[s['name']]=s
        return s
    for item in schema:
        if 'csv' in item:add(item,source_csv(root,item['csv'],item['headers']))
        elif 'static_rows' in item:
            rows=[list(x) for x in item['static_rows']]
            add(item,rows)
        elif item['name']=='Table1_Efficacy':
            raw=lookup['Morph_48_Cells']['values'][4:];h=lookup['Morph_48_Cells']['values'][3]
            rr=[dict(zip(h,r)) for r in raw]
            control={(r['Isolate'],r['UV_mJ_cm2']):i+5 for i,r in enumerate(rr) if r['Chemical']=='Control'}
            rows=[];forms=[]
            for i,r in enumerate(rr,5):
                ref=control[(r['Isolate'],r['UV_mJ_cm2'])];cm=float(rr[ref-5]['mean_area_mm2'])
                estimate=100*(1-float(r['mean_area_mm2'])/cm)
                rows.append([r['Isolate'],'EtOH' if r['Chemical']=='Control' else r['Chemical'],r['UV_mJ_cm2'],r['n'],r['mean_area_mm2'],r['area_ci_low'],r['area_ci_high'],cm,estimate,None if r['Chemical']=='Control' else r['inhibition_ci_low'],None if r['Chemical']=='Control' else r['inhibition_ci_high']])
                forms += [(f'H{i}',f'=E{ref}'),(f'I{i}',f'=100*(1-E{i}/H{i})')]
                if abs(estimate-float(r['inhibition_percent']))>1e-9:raise ArithmeticError('Table1 display calculation mismatch')
            add(item,rows,forms)
        elif item['name']=='Table2_HSI':
            et=lookup['HSI_Target_Effects'];st=lookup['HSI_Band_Tests']
            e=[dict(zip(et['values'][3],r)) for r in et['values'][4:]]
            sp={(r['comparison'],r['wavelength_nm']):r for r in [dict(zip(st['values'][3],q)) for q in st['values'][4:]]}
            rows=[];forms=[]
            for i,r in enumerate(e,5):
                q=sp[(r['comparison'],r['wavelength_nm'])]
                estimate=r['mean_UV70']-r['mean_UV0']
                rows.append([r['condition_from_codebook'],r['wavelength_nm'],r['n0'],r['n70'],r['mean_UV0'],r['mean_UV70'],estimate,r['ci_low'],r['ci_high'],q['p_raw_difference'],q['p_max_abs_difference'],q['p_max_abs_t'],q['legacy_mc_max_abs_diff_p'],'Original target' if r['wavelength_nm']==518.8 else 'Same-data scan-selected endpoint'])
                forms.append((f'G{i}',f'=F{i}-E{i}'))
                if abs(estimate-r['mean_difference_observed'])>1e-9:raise ArithmeticError('Table2 display calculation mismatch')
            add(item,rows,forms)
        elif item['name']=='NFI_Map_Source':add(item,source_csv(root,'NFI_Map_Source.csv',item['headers']))
        elif item['name']=='Workbook_Checks':
            t=lookup['Table1_Efficacy']['values'][4:];raw=lookup['Morph_48_Cells']['values'][4:];hh=lookup['Morph_48_Cells']['values'][3]
            rows=[];forms=[]
            for i,(r,rr) in enumerate(zip(t,raw),5):
                d=dict(zip(hh,rr));observed=r[8];expected=d['inhibition_percent']
                rows.append([f"Table1 inhibition / {d['Isolate']}/{d['Chemical']}/{d['UV_mJ_cm2']}",expected,observed,abs(expected-observed)])
                forms.extend([(f'C{i}',f"='Table1_Efficacy'!I{i}"),(f'D{i}',f'=ABS(B{i}-C{i})')])
            add(item,rows,forms)
        elif item['name']=='INDEX':
            rows=[[s['name'],s['data_rows'],s['columns'],('Metadata' if s['name'] in ['README','CODEBOOK','ANALYSIS_PARAMETERS'] else 'Analysis / provenance'),s['values'][1][0],s['source']] for s in sheets]
            add(item,rows)
    return {'schema_version':'1.2.0','sheets':sheets}


def cached_formulas(path,snapshot):
    """Cache only generated display formulas checked against their source values.
    This is not a general Excel calculation engine. Excel recalculation is enabled.
    """
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    def index(addr):
        m=re.fullmatch(r'([A-Z]+)([0-9]+)',addr);n=0
        for c in m[1]:n=n*26+ord(c)-64
        return int(m[2])-1,n-1
    bysheet={s['name']:s for s in snapshot['sheets']}
    for sheet in snapshot['sheets']:
        for addr,formula in sheet['formulas']:
            row,col=index(addr);expected=sheet['values'][row][col]
            if formula.startswith("='"):
                m=re.fullmatch(r"='([^']+)'!([A-Z]+[0-9]+)",formula)
                r,c=index(m[2]);actual=bysheet[m[1]]['values'][r][c]
            else:
                expr=formula[1:]
                def repl(m):
                    r,c=index(m[0]);v=sheet['values'][r][c]
                    if not isinstance(v,(int,float)):raise ValueError('Non-numeric formula reference')
                    return repr(v)
                expr=re.sub(r'\b[A-Z]+[0-9]+\b',repl,expr).replace('ABS','abs')
                # Generated grammar only: numbers, operators, parentheses and abs.
                if not re.fullmatch(r'[0-9eE+\-*/(). abs]+',expr):raise ValueError('Unexpected display formula')
                actual=eval(expr,{'__builtins__':{}},{'abs':abs})
            if not math.isfinite(float(actual)) or abs(actual-expected)>1e-9:
                raise ArithmeticError('Cached formula mismatch '+sheet['name']+'!'+addr)
    tmp=path.with_name(path.stem+'.cache.tmp.xlsx')
    with zipfile.ZipFile(path) as zin,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=zin.read(info.filename)
            m=re.fullmatch(r'xl/worksheets/sheet([0-9]+)\.xml',info.filename)
            if m:
                spec=snapshot['sheets'][int(m[1])-1];formulas=dict(spec['formulas'])
                if formulas:
                    rt=ET.fromstring(data)
                    for c in rt.findall('.//s:sheetData/s:row/s:c',ns):
                        addr=c.attrib.get('r')
                        if addr not in formulas:continue
                        r,k=index(addr);v=c.find('s:v',ns)
                        if v is None:v=ET.SubElement(c,'{'+ns['s']+'}v')
                        v.text=repr(spec['values'][r][k])
                    data=ET.tostring(rt,encoding='utf-8',xml_declaration=True)
            zout.writestr(info,data)
    # Check the serialized XML bytes, including the encoding declaration.
    with zipfile.ZipFile(tmp) as check:
        for member in check.namelist():
            if member.endswith('.xml'):ET.fromstring(check.read(member))
    tmp.replace(path)


def write_portable(snapshot,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();wb.remove(wb.active)
    wb.properties.creator='';wb.properties.lastModifiedBy=''
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    for s in snapshot['sheets']:
        ws=wb.create_sheet(s['name']);n=s['columns'];end=s['data_rows']+4
        for r in s['values']:ws.append(r)
        for row in ws:
            for c in row:
                c.font=Font(name='Calibri',size=10)
                if isinstance(c.value,str) and c.value.startswith('='):c.data_type='s'
                if isinstance(c.value,float):c.number_format='0.000000'
        for addr,f in s['formulas']:ws[addr]=f
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
        ws['A1'].font=Font(name='Calibri',bold=True,color='FFFFFF',size=13)
        ws['A1'].fill=PatternFill('solid',fgColor='17365D');ws.row_dimensions[1].height=28
        ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=44
        for c in ws[4]:
            c.font=Font(name='Calibri',bold=True,color='17365D',size=10);c.fill=PatternFill('solid',fgColor='DCE6F1');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[4].height=48;ws.freeze_panes='A5'
        for j,h in enumerate(s['values'][3],1):
            letter=column(j);ws.column_dimensions[letter].width=24 if j==1 else 18
            if re.search('note|detail|scope|source|file|method|coding|comparison',str(h),re.I):ws.column_dimensions[letter].width=36
            if re.search(r'\bp_|p.value|p_holm|p_max|p_exact|PR\(>F\)',str(h),re.I):
                for row in range(5,end+1):ws.cell(row,j).number_format='0.000E+00'
        if s['data_rows']:
            tab=Table(displayName='T_'+s['name'],ref=f'A4:{column(n)}{end}')
            tab.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True)
            ws.add_table(tab)
        if s['name']=='README':
            ws.column_dimensions['B'].width=110
            for row in range(5,end+1):ws.row_dimensions[row].height=38;ws.cell(row,2).alignment=Alignment(wrap_text=True,vertical='center')
        if s['name']=='INDEX':
            for c in ['E','F']:ws.column_dimensions[c].width=55
            for row in range(5,end+1):
                ws.row_dimensions[row].height=48
                for c in ws[row]:c.alignment=Alignment(wrap_text=True,vertical='center')
        if s['name']=='Table1_Efficacy':
            for row in ws.iter_rows(min_row=5,max_col=11,min_col=5):
                for c in row:c.number_format='0.00'
    wb.save(path);cached_formulas(path,snapshot)


def build_workbook(root,path):
    root=Path(root);path=Path(path)
    if path.exists():raise FileExistsError('Refusing to overwrite workbook: '+str(path))
    spec=make_snapshot(root)
    path.parent.mkdir(parents=True,exist_ok=True)
    write_portable(spec,path)
    (path.parent/'workbook_snapshot.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf8')
    (path.parent/'workbook_validation.json').write_text(json.dumps({'status':'PASS','sheets':len(spec['sheets']),'display_formulas':sum(len(s['formulas']) for s in spec['sheets']),'formula_tolerance':1e-9,'statistics_recomputed_by_exporter':False},indent=2),encoding='utf8')
    return path
